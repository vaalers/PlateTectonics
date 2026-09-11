#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from __future__ import annotations

import argparse
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import List, Optional

from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt

HEATMAP_NAME_RE = re.compile(
    r"^(?P<date>\d{6}|\d{4})_responder_annotation_heatmap_plate(?P<plate>\d+)"
    r"_experiment(?P<experiments>[\d-]+|NA)_eval(?P<eval>[\d-]+|NA)\.png$",
    re.IGNORECASE,
)

SLIDE_WIDTH_IN = 13.333
SLIDE_HEIGHT_IN = 7.5
MARGIN_IN = 0.4
TITLE_HEIGHT_IN = 0.6


@dataclass(frozen=True)
class HeatmapImage:
    path: Path
    date_code: str
    plate: int
    experiments: str
    eval_label: str

    @property
    def parsed_date(self) -> Optional[datetime]:
        if len(self.date_code) == 6:
            try:
                return datetime.strptime(self.date_code, "%m%d%y")
            except ValueError:
                return None
        if len(self.date_code) == 4:
            # Year-only rollup file (no specific date) - sort to the end of that year.
            try:
                return datetime(int(self.date_code), 12, 31)
            except ValueError:
                return None
        return None

    @property
    def sort_key(self):
        dt = self.parsed_date or datetime.max
        return (dt, self.plate, self.experiments)

    @property
    def title(self) -> str:
        if len(self.date_code) == 6 and self.parsed_date is not None:
            display_date = self.parsed_date.strftime("%m/%d/%Y")
        elif len(self.date_code) == 4:
            display_date = f"Year {self.date_code} (rollup)"
        else:
            display_date = self.date_code
        return (
            f"{display_date}  \u2014  Plate {self.plate}  "
            f"\u2014  Experiment(s) {self.experiments}  \u2014  Eval {self.eval_label}"
        )


def find_heatmap_images(root: Path) -> List[HeatmapImage]:
    images: List[HeatmapImage] = []
    for path in root.rglob("*_responder_annotation_heatmap_plate*_experiment*_eval*.png"):
        m = HEATMAP_NAME_RE.match(path.name)
        if not m:
            continue
        images.append(
            HeatmapImage(
                path=path,
                date_code=m.group("date"),
                plate=int(m.group("plate")),
                experiments=m.group("experiments"),
                eval_label=m.group("eval"),
            )
        )
    images.sort(key=lambda img: img.sort_key)
    return images


def add_title(slide, text: str) -> None:
    box = slide.shapes.add_textbox(
        Inches(MARGIN_IN), Inches(0.15), Inches(SLIDE_WIDTH_IN - 2 * MARGIN_IN), Inches(TITLE_HEIGHT_IN)
    )
    tf = box.text_frame
    tf.word_wrap = True
    run = tf.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(20)
    run.font.bold = True


def add_scaled_picture(slide, image: HeatmapImage) -> None:
    with Image.open(image.path) as im:
        px_width, px_height = im.size

    area_left = MARGIN_IN
    area_top = TITLE_HEIGHT_IN + 0.3
    area_width = SLIDE_WIDTH_IN - 2 * MARGIN_IN
    area_height = SLIDE_HEIGHT_IN - area_top - MARGIN_IN

    aspect = px_width / px_height
    width_in = area_width
    height_in = width_in / aspect
    if height_in > area_height:
        height_in = area_height
        width_in = height_in * aspect

    left_in = area_left + (area_width - width_in) / 2
    top_in = area_top + (area_height - height_in) / 2

    slide.shapes.add_picture(
        str(image.path),
        Inches(left_in),
        Inches(top_in),
        width=Inches(width_in),
        height=Inches(height_in),
    )


def build_presentation(images: List[HeatmapImage]) -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_WIDTH_IN)
    prs.slide_height = Inches(SLIDE_HEIGHT_IN)
    blank_layout = prs.slide_layouts[6]

    for image in images:
        slide = prs.slides.add_slide(blank_layout)
        add_title(slide, image.title)
        add_scaled_picture(slide, image)

    return prs


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Collect responder annotation heatmap PNGs from every experiment "
        "under an analyzed data root into a single PowerPoint, one image per slide."
    )
    parser.add_argument(
        "--root",
        required=True,
        help="Root directory to search recursively for heatmap PNGs.",
    )
    parser.add_argument(
        "--output",
        default="responder_annotation_heatmaps.pptx",
        help="Output .pptx path.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    root = Path(args.root)
    output = Path(args.output)

    print(f"[INFO] Searching for heatmap PNGs under: {root}")
    images = find_heatmap_images(root)
    print(f"[INFO] Found {len(images)} heatmap image(s)")
    if not images:
        print("No heatmap images found; nothing to build.")
        return

    prs = build_presentation(images)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    print(f"Wrote {len(images)} slide(s) to {output}")


if __name__ == "__main__":
    main()
